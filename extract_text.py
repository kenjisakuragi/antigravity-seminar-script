# -*- coding: utf-8 -*-
import sys
import os

try:
    from pptx import Presentation
except ImportError:
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-pptx"])
    from pptx import Presentation

pptx_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "AI副業セミナー_v3.pptx")
prs = Presentation(pptx_path)

output_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "slides_text.txt")

with open(output_path, "w", encoding="utf-8") as f:
    for i, slide in enumerate(prs.slides, 1):
        f.write(f"=== スライド {i} ===\n")
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = para.text.strip()
                    if text:
                        f.write(text + "\n")
            if shape.has_table:
                table = shape.table
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    f.write(" | ".join(cells) + "\n")
        f.write("\n")

print(f"Done! Output saved to {output_path}")
