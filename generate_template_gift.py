# -*- coding: utf-8 -*-
"""
第2回 プレゼント用テンプレート集をWordで出力
「台本生成プロンプト」「背景画像プロンプト」「キャラクタープロンプト」3点セット
"""
import sys, os

try:
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
except ImportError:
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx"])
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUT_PATH = os.path.join(SCRIPT_DIR, "第2回_テンプレート3点セット.docx")

# ── ヘルパー ──────────────────────────────────────
def set_font(run, size_pt, bold=False, color=None, name="メイリオ"):
    run.font.name = name
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)

def heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    colors = {1: (0x05, 0x96, 0x69), 2: (0x05, 0x96, 0x69), 3: (0x03, 0x7A, 0x52)}
    sizes  = {1: 20, 2: 15, 3: 13}
    set_font(run, sizes[level], bold=True, color=colors[level])
    return p

def body(doc, text, indent=0, size=11, color=None, bold=False):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(indent)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text)
    set_font(run, size, bold=bold, color=color)
    return p

def box(doc, lines, bg_rgb=(0xF0, 0xFD, 0xF4), border_rgb=(0x05, 0x96, 0x69)):
    """枠付きボックス（シェーディングで代用）"""
    for i, line in enumerate(lines):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent  = Cm(0.8)
        p.paragraph_format.right_indent = Cm(0.8)
        p.paragraph_format.space_after  = Pt(1)
        if i == 0:
            p.paragraph_format.space_before = Pt(6)
        if i == len(lines) - 1:
            p.paragraph_format.space_after = Pt(6)
        run = p.add_run(line)
        set_font(run, 11, color=(0x02, 0x4B, 0x2E) if line.startswith("【") else None)
        # シェーディング
        pPr = p._p.get_or_add_pPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        r, g, b = bg_rgb
        shd.set(qn('w:fill'), f'{r:02X}{g:02X}{b:02X}')
        pPr.append(shd)
    return

def example_box(doc, label, lines):
    """例示ボックス（薄紫背景）"""
    p0 = doc.add_paragraph()
    p0.paragraph_format.space_before = Pt(4)
    p0.paragraph_format.space_after  = Pt(2)
    r0 = p0.add_run(f"✏️ 入力例：{label}")
    set_font(r0, 11, bold=True, color=(0x59, 0x39, 0xAB))
    for i, line in enumerate(lines):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent  = Cm(0.8)
        p.paragraph_format.right_indent = Cm(0.8)
        p.paragraph_format.space_after  = Pt(1)
        if i == len(lines) - 1:
            p.paragraph_format.space_after = Pt(8)
        run = p.add_run(line)
        set_font(run, 10.5)
        pPr = p._p.get_or_add_pPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), 'F5F3FF')
        pPr.append(shd)

def tip(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run(f"📌 {text}")
    set_font(run, 10.5, bold=True, color=(0xD9, 0x77, 0x06))

def divider(doc):
    doc.add_paragraph("─" * 55)

# ── 本文 ──────────────────────────────────────────
doc = Document()
sec = doc.sections[0]
sec.left_margin   = Cm(2.0)
sec.right_margin  = Cm(2.0)
sec.top_margin    = Cm(2.0)
sec.bottom_margin = Cm(2.0)

style = doc.styles['Normal']
style.font.name = 'メイリオ'
style.font.size = Pt(11)
style.paragraph_format.line_spacing = 1.4

# ══ 表紙 ══════════════════════════════════════════
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(40)
r = p.add_run("らくらくAI副業キャンパス 第2回\n配布テンプレート")
set_font(r, 18, bold=True, color=(0x05, 0x96, 0x69))

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("プロンプト 3点セット")
set_font(r, 26, bold=True, color=(0x06, 0x4E, 0x3B))

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("① 台本生成プロンプト　② 背景画像プロンプト　③ キャラクタープロンプト")
set_font(r, 13, color=(0x37, 0x47, 0x51))

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("\n「〇〇」の部分を変えるだけで、何本でも作れます！")
set_font(r, 12, bold=True, color=(0xD9, 0x77, 0x06))

doc.add_page_break()

# ══ 使い方ガイド ══════════════════════════════════
heading(doc, "📖 このテンプレートの使い方", level=1)
body(doc, "1. Gemini（gemini.google.com）または画像生成ツールを開く")
body(doc, "2. 使いたいプロンプトをそのままコピーする")
body(doc, "3.「〇〇」の部分を自分のジャンル・テーマ・好みに書き換える")
body(doc, "4. Enterキーを押すだけ！完成です")
doc.add_paragraph()
body(doc, "▶ 一度作ったプロンプトは保存しておくと、次回からさらに楽になります。", bold=True)
body(doc, "▶ 「なんか違う」と思ったら、Geminiに「もう少し優しく書いて」と追加で伝えてみてください。")

divider(doc)
doc.add_paragraph()

# ══════════════════════════════════════════════════
# ① 台本生成プロンプト
# ══════════════════════════════════════════════════
heading(doc, "① 台本生成プロンプト", level=1)
body(doc, "Gemini（gemini.google.com）に貼り付けて使います", color=(0x55, 0x65, 0x74))

heading(doc, "▼ 基本テンプレート（コピペ用）", level=2)
box(doc, [
    "以下の条件でYouTube動画の台本を書いてください。",
    "",
    "【ジャンル】〇〇（例：スピリチュアル / 猫の日常 / 雑学 / 癒やし）",
    "【テーマ】〇〇（例：運気が上がる朝の習慣 / 猫が甘えてくる理由）",
    "【台本の型】リスト系（「〇〇な人の5つの特徴」のような構成）",
    "【文字数】2,000字程度",
    "【口調】落ち着いた、優しい語り口調。難しい言葉は使わない",
    "【構成】",
    "  ① 冒頭のつかみ（最初の30秒で興味を引く）",
    "  ② 本編（3〜5項目のリスト形式）",
    "  ③ まとめ（温かい言葉で締める）",
    "  ④ 次も見たくなる一言（チャンネル登録を促す）",
])

heading(doc, "▼ ジャンル別 入力例", level=2)

example_box(doc, "スピリチュアル系", [
    "【ジャンル】スピリチュアル",
    "【テーマ】直感を磨くための3つの習慣",
    "【台本の型】リスト系",
    "【文字数】2,000字程度",
    "【口調】神秘的で落ち着いた、語りかけるような口調",
    "【構成】冒頭のつかみ → 本編（3項目）→ まとめ → 次回予告",
])

example_box(doc, "猫・ペット系", [
    "【ジャンル】猫の日常",
    "【テーマ】猫が「ふみふみ」をする本当の理由",
    "【台本の型】教養系（知識を優しく解説）",
    "【文字数】1,500字程度",
    "【口調】温かく、猫好きに語りかけるような優しい口調",
    "【構成】冒頭のつかみ → 解説（理由3つ）→ まとめ",
])

example_box(doc, "雑学・教養系", [
    "【ジャンル】日本の歴史雑学",
    "【テーマ】知らないと恥ずかしい江戸時代の常識5選",
    "【台本の型】リスト系",
    "【文字数】2,500字程度",
    "【口調】わかりやすく、親しみやすい語り口",
    "【構成】つかみ → 5項目リスト → まとめ",
])

tip(doc, "「なんか違う」と思ったら→「もう少し感動的な表現にして」「もっと短くして」と追加で伝えればOK！")

heading(doc, "▼ 台本のクオリティをさらに上げる追加指示", level=2)
body(doc, "・「最初の3行で視聴者の心をつかむようにして」")
body(doc, "・「体験談や具体的なエピソードを入れて」")
body(doc, "・「読んで思わず涙が出るような感動的な締めにして」")
body(doc, "・「小学生でもわかるような言葉で書き直して」")

doc.add_paragraph()
divider(doc)
doc.add_paragraph()

# ══════════════════════════════════════════════════
# ② 背景画像プロンプト
# ══════════════════════════════════════════════════
heading(doc, "② 背景画像プロンプト", level=1)
body(doc, "Geminiの画像生成、またはCanva AIに貼り付けて使います", color=(0x55, 0x65, 0x74))

heading(doc, "▼ 基本テンプレート（コピペ用）", level=2)
box(doc, [
    "YouTube動画の背景画像を作ってください。",
    "",
    "【テーマ・世界観】〇〇（例：神秘的な森 / 温かいカフェ / 宇宙）",
    "【スタイル】〇〇（例：水彩画風 / アニメ風 / ファンタジーイラスト / 写真風）",
    "【雰囲気・色】〇〇（例：落ち着いた紫と金 / 温かいオレンジ / 幻想的な青）",
    "【サイズ】横長（16:9）",
    "【注意】文字は一切入れないでください。キャラクターも入れないでください。",
    "        背景だけのシーンにしてください。",
])

heading(doc, "▼ ジャンル別 入力例", level=2)

example_box(doc, "スピリチュアル系", [
    "【テーマ・世界観】神秘的な森の中。月の光が差し込んでいる。",
    "【スタイル】幻想的なデジタルイラスト風",
    "【雰囲気・色】深い紫と金色。神秘的で美しい雰囲気",
    "【サイズ】横長（16:9）",
    "【注意】文字なし、キャラクターなし。背景のみ。",
])

example_box(doc, "猫・ほっこり系", [
    "【テーマ・世界観】温かみのある和風の縁側。日差しが差し込んでいる。",
    "【スタイル】水彩画風のほっこりイラスト",
    "【雰囲気・色】温かいオレンジと茶色。柔らかい光",
    "【サイズ】横長（16:9）",
    "【注意】文字なし、キャラクターなし。",
])

example_box(doc, "雑学・教養系", [
    "【テーマ・世界観】本がたくさん並ぶ落ち着いた図書館。",
    "【スタイル】シンプルなイラスト風",
    "【雰囲気・色】ダークブラウンと温かいアンバー。知的な雰囲気",
    "【サイズ】横長（16:9）",
    "【注意】文字なし、キャラクターなし。",
])

tip(doc, "背景は「シンプルすぎるくらい」がちょうどいい。ごちゃごちゃしているとキャラクターが目立たなくなります。")

heading(doc, "▼ スタイルの参考一覧", level=2)
body(doc, "水彩画風 ／ アニメ風 ／ 油絵風 ／ ファンタジーイラスト")
body(doc, "ミニマルデザイン ／ 和風イラスト ／ ポップアート風 ／ 写実的な風景画")
body(doc, "サイバーパンク風 ／ 星空・宇宙 ／ 森林浴 ／ 日本の四季")

doc.add_paragraph()
divider(doc)
doc.add_paragraph()

# ══════════════════════════════════════════════════
# ③ キャラクタープロンプト
# ══════════════════════════════════════════════════
heading(doc, "③ キャラクタープロンプト", level=1)
body(doc, "Geminiの画像生成、またはCanva AIに貼り付けて使います", color=(0x55, 0x65, 0x74))

heading(doc, "▼ 基本テンプレート（コピペ用）", level=2)
box(doc, [
    "動画のナレーターとして使うオリジナルキャラクターを作ってください。",
    "",
    "【キャラクターの雰囲気】〇〇（例：神秘的 / 優しい / 元気 / 知的 / ほっこり）",
    "【性別・年齢イメージ】〇〇（例：20代女性 / 30代男性 / 性別不明の妖精）",
    "【スタイル】〇〇（例：アニメ風 / ゆるいイラスト風 / リアル系）",
    "【衣装・特徴】〇〇（例：白いワンピース / 和服 / ローブ / カジュアルな服）",
    "【髪の色・スタイル】〇〇（例：長い黒髪 / 銀髪ショート / 天然パーマの茶髪）",
    "【表情】優しい微笑み（最初はこれだけでOK）",
    "【背景】白背景（または透過背景）",
    "【注意】上半身〜全身が映る構図にしてください。",
])

heading(doc, "▼ ジャンル別 入力例", level=2)

example_box(doc, "スピリチュアル系キャラクター", [
    "【キャラクターの雰囲気】神秘的で落ち着いた、霊的な存在感",
    "【性別・年齢イメージ】20代後半の女性",
    "【スタイル】アニメ風イラスト",
    "【衣装・特徴】深紫のローブ。胸元に三日月のアクセサリー",
    "【髪の色・スタイル】長い銀髪。緩やかなウェーブ",
    "【表情】穏やかで、少し微笑んでいる",
    "【背景】白背景",
    "【注意】上半身が映る構図。威圧的にならず、親しみやすく。",
])

example_box(doc, "猫・ほっこり系キャラクター", [
    "【キャラクターの雰囲気】ゆるくて温かい、癒やし系",
    "【性別・年齢イメージ】猫の耳がついた女の子（年齢不詳）",
    "【スタイル】ゆるいイラスト風（シンプルなデザイン）",
    "【衣装・特徴】白いエプロンドレス。ふわふわした猫耳と尻尾",
    "【髪の色・スタイル】茶色のショートヘア",
    "【表情】にこにことした笑顔",
    "【背景】白背景",
    "【注意】全身が見える構図。丸みのあるデザインで。",
])

example_box(doc, "雑学・教養系キャラクター", [
    "【キャラクターの雰囲気】知的で親しみやすい、先生のような雰囲気",
    "【性別・年齢イメージ】30代男性",
    "【スタイル】シンプルなアニメ風",
    "【衣装・特徴】白いシャツ＋カジュアルなジャケット。丸メガネ",
    "【髪の色・スタイル】黒髪のナチュラルスタイル",
    "【表情】優しく語りかけるような表情",
    "【背景】白背景",
    "【注意】上半身〜腰くらいまで映る構図。",
])

tip(doc, "まず1体作ってみて、気に入らない部分だけ「髪をもっと長くして」「衣装を〇〇に変えて」と調整しましょう！")

heading(doc, "▼ 世界観を合わせるコツ", level=2)
body(doc, "背景画像とキャラクターは「同じ雰囲気の言葉」を使うと世界観が統一されます。")
doc.add_paragraph()

t = doc.add_table(rows=4, cols=3)
t.style = 'Light Grid Accent 6'
headers = ["ジャンル", "背景の言葉", "キャラクターの言葉"]
for i, h in enumerate(headers):
    c = t.rows[0].cells[i]
    c.text = h
    for p in c.paragraphs:
        for r in p.runs:
            r.font.bold = True
            r.font.name = 'メイリオ'
rows_data = [
    ["スピリチュアル", "神秘的・幻想的・月・紫・金色", "神秘的・霊的・ローブ・銀髪"],
    ["猫・癒やし", "温かい・縁側・日差し・オレンジ", "ゆるい・猫耳・エプロン・茶色"],
    ["雑学・教養", "図書館・本・ダークブラウン・知的", "メガネ・ジャケット・知的・親しみやすい"],
]
for ri, row_data in enumerate(rows_data, 1):
    for ci, val in enumerate(row_data):
        c = t.rows[ri].cells[ci]
        c.text = val
        for p in c.paragraphs:
            for r in p.runs:
                r.font.name = 'メイリオ'
                r.font.size = Pt(10)

doc.add_paragraph()
divider(doc)
doc.add_paragraph()

# ══ まとめ・使い方ガイド ══
heading(doc, "📋 まとめ｜3点セットの活用フロー", level=1)
body(doc, "STEP 1：台本プロンプト → Geminiで台本を生成", bold=True)
body(doc, "        台本が完成したら、タイトル候補も5個作ってもらおう", indent=1)
doc.add_paragraph()
body(doc, "STEP 2：背景画像プロンプト → 動画の「舞台」を作る", bold=True)
body(doc, "        3〜5パターン生成して、一番気に入ったものを選ぼう", indent=1)
doc.add_paragraph()
body(doc, "STEP 3：キャラクタープロンプト → あなたの「分身」を誕生させる", bold=True)
body(doc, "        来週、このキャラクターに声を入れてリップシンク動画を作ります！", indent=1)
doc.add_paragraph()
body(doc, "📌 全部できたら、タイトルをDiscordに投稿して仲間に宣言しましょう！",
     bold=True, color=(0xD9, 0x77, 0x06))

doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("らくらくAI副業キャンパス | 第2回配布資料")
set_font(r, 10, color=(0x9C, 0xA3, 0xAF))

doc.save(OUT_PATH)
print(f"完了！ → {OUT_PATH}")
