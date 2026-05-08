# -*- coding: utf-8 -*-
"""Extract all slide text content from the complete PPTX"""
import sys, os
sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "Google_Slide_generator"))
from pptx import Presentation

prs = Presentation(os.path.join(os.path.dirname(os.path.abspath(__file__)), "AI副業セミナー_完全版.pptx"))
out_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "slide_content_dump.txt")

with open(out_path, "w", encoding="utf-8") as f:
    for i, slide in enumerate(prs.slides, 1):
        f.write(f"=== スライド {i} ===\n")
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    t = para.text.strip()
                    if t:
                        f.write(f"  {t}\n")
        f.write("\n")

print(f"完了: {out_path}")
print(f"スライド数: {len(prs.slides)}")
