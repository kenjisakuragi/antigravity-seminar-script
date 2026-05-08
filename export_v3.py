# -*- coding: utf-8 -*-
"""Export seminar_script_v3.md to Word"""
import sys, os, re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
MD_PATH = os.path.join(SCRIPT_DIR, "seminar_script_v3.md")
OUT_PATH = os.path.join(SCRIPT_DIR, "セミナー読み原稿_v3.docx")

with open(MD_PATH, "r", encoding="utf-8") as f:
    lines = f.readlines()
print(f"Read {len(lines)} lines")

doc = Document()
style = doc.styles["Normal"]
style.font.name = "メイリオ"
style.font.size = Pt(11)
style.paragraph_format.line_spacing = 1.5

for i in range(1, 4):
    h = doc.styles[f"Heading {i}"]
    h.font.name = "メイリオ"
    h.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)

# Cover
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("\n\n\n").font.size = Pt(14)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("AI副業セミナー")
r.font.size = Pt(28)
r.font.bold = True
r.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("読み原稿（完全版 v3）")
r.font.size = Pt(22)
r.font.bold = True
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("\n対応：AI副業セミナー_完全版_修正版.pptx（全79枚）\n執筆方針：17原則に準拠")
r.font.size = Pt(12)
r.font.color.rgb = RGBColor(0x60, 0x60, 0x60)
doc.add_page_break()

# Legend
doc.add_heading("記号凡例", level=1)
items = [
    ("【間 ○秒】", "指定秒数の沈黙"),
    ("【スライド切替】", "次スライドへ"),
    ("【実演開始】〜【実演終了】", "デモ区間"),
    ("【声のトーン：↑/↓】", "テンション指示"),
    ("【ゆっくり】", "重要、ゆっくり読む"),
    ("📌", "特に強調"),
    ("※", "講師用メモ（読まない）"),
]
table = doc.add_table(rows=len(items) + 1, cols=2)
table.style = "Light Grid Accent 1"
table.rows[0].cells[0].text = "記号"
table.rows[0].cells[1].text = "意味"
for cell in table.rows[0].cells:
    for para in cell.paragraphs:
        for run in para.runs:
            run.font.bold = True
for idx, (sym, desc) in enumerate(items, 1):
    table.rows[idx].cells[0].text = sym
    table.rows[idx].cells[1].text = desc
doc.add_page_break()


def add_styled(doc, text, stage=False):
    p = doc.add_paragraph()
    if stage:
        p.paragraph_format.left_indent = Cm(1)
        r = p.add_run(text)
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0, 0x80, 0)
        r.font.italic = True
        return
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            r = p.add_run(part[2:-2])
            r.font.bold = True
        else:
            p.add_run(part)


for line in lines:
    line = line.rstrip("\r\n")
    if not line.strip():
        continue
    if line.strip() == "---":
        continue
    if line.startswith("# ━"):
        continue
    stripped = line.strip()
    if stripped.startswith("```"):
        continue
    # Section headings for page break
    if line.startswith("# 第") or line.startswith("# 補足"):
        doc.add_page_break()
        title = line.lstrip("# ").strip()
        doc.add_heading(title, level=1)
        continue
    # Skip other top-level comments
    if line.startswith("# "):
        continue
    # H2
    if line.startswith("## "):
        title = line.lstrip("# ").strip()
        doc.add_heading(title, level=2)
        continue
    # Blockquote
    if line.startswith("> "):
        p = doc.add_paragraph(line[2:].strip())
        p.paragraph_format.left_indent = Cm(1.5)
        if p.runs:
            p.runs[0].font.italic = True
            p.runs[0].font.color.rgb = RGBColor(0x60, 0x60, 0x60)
        continue
    # Table rows
    if stripped.startswith("|"):
        add_styled(doc, stripped)
        continue
    # Stage directions
    stage_keys = ["スライド切替", "間 ", "間：", "声のトーン", "ゆっくり", "実演", "カウントダウン", "リアルタイム"]
    if stripped.startswith("【") and any(k in stripped for k in stage_keys):
        add_styled(doc, stripped, stage=True)
        continue
    # Notes
    if stripped.startswith("※") or stripped.startswith("→ "):
        add_styled(doc, stripped, stage=True)
        continue
    # Normal text
    add_styled(doc, stripped)

doc.save(OUT_PATH)
print(f"Saved: {OUT_PATH}")
print(f"Size: {os.path.getsize(OUT_PATH)} bytes")
