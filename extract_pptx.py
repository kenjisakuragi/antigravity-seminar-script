import sys, os
try:
    from pptx import Presentation
except ImportError:
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-pptx"])
    from pptx import Presentation

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
pptx_path = os.path.join(SCRIPT_DIR, "講義03_AI動画②_v3.pptx")
prs = Presentation(pptx_path)
out = os.path.join(SCRIPT_DIR, "講義03_v3_slides.txt")

with open(out, "w", encoding="utf-8") as f:
    for i, slide in enumerate(prs.slides, 1):
        f.write(f"=== スライド {i} ===\n")
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    t = para.text.strip()
                    if t:
                        f.write(t + "\n")
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    f.write(" | ".join(c.text.strip() for c in row.cells) + "\n")
        f.write("\n")

print(f"完了 → {out}  ({len(prs.slides)}枚)")
