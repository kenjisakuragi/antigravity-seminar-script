# -*- coding: utf-8 -*-
"""
サブスク2プラン（Basic 5,000円 / Premium 10,000円）受講者案内スライド

CLAUDE.md デザイン仕様準拠：
- フォント：メイリオ
- 見出し色：RGB(14, 165, 233)  ライトブルー
- 強調色  ：RGB(234, 88, 12)   オレンジ
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

FONT = "メイリオ"
BLUE = RGBColor(0x0E, 0xA5, 0xE9)
BLUE_DARK = RGBColor(0x02, 0x84, 0xC7)
ORANGE = RGBColor(0xEA, 0x58, 0x0C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x0F, 0x17, 0x2A)
GRAY = RGBColor(0x64, 0x74, 0x8B)
LIGHT_BG = RGBColor(0xF1, 0xF5, 0xF9)
PREMIUM_GOLD = RGBColor(0xCA, 0x8A, 0x04)


def set_run(run, text, size=18, color=DARK, bold=False, font=FONT):
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_text(slide, left, top, width, height, text, *, size=18, color=DARK,
             bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    p = tf.paragraphs[0]
    p.alignment = align
    set_run(p.add_run(), text, size=size, color=color, bold=bold)
    return tb


def add_rect(slide, left, top, width, height, fill, line=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(1)
    shp.shadow.inherit = False
    return shp


def add_round_rect(slide, left, top, width, height, fill, line=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shp.adjustments[0] = 0.10
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(1.25)
    shp.shadow.inherit = False
    return shp


def slide_blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def add_header_bar(slide, title, subtitle=None):
    add_rect(slide, Inches(0), Inches(0), Inches(13.333), Inches(0.55), BLUE)
    add_text(slide, Inches(0.4), Inches(0.05), Inches(12.5), Inches(0.45),
             "らくらくAI副業キャンパス｜サブスクリプションのご案内",
             size=14, color=WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, Inches(0.5), Inches(0.75), Inches(12.3), Inches(0.7),
             title, size=30, color=DARK, bold=True)
    if subtitle:
        add_text(slide, Inches(0.5), Inches(1.45), Inches(12.3), Inches(0.4),
                 subtitle, size=14, color=GRAY)
    add_rect(slide, Inches(0.5), Inches(1.85), Inches(0.6), Emu(38000), ORANGE)


# ─────────────────────────────────────────────
# Slide 1: Title
# ─────────────────────────────────────────────
def slide_title(prs):
    s = slide_blank(prs)
    add_rect(s, 0, 0, prs.slide_width, prs.slide_height, BLUE)
    add_rect(s, 0, Inches(5.2), prs.slide_width, Inches(2.3), BLUE_DARK)
    add_text(s, Inches(0.6), Inches(1.0), Inches(12.1), Inches(0.5),
             "らくらくAI副業キャンパス", size=18, color=WHITE, bold=True)
    add_text(s, Inches(0.6), Inches(1.7), Inches(12.1), Inches(2.0),
             "サブスクリプション\n受講後も学びと収益を止めない仕組み",
             size=44, color=WHITE, bold=True)
    add_text(s, Inches(0.6), Inches(5.45), Inches(12.1), Inches(0.5),
             "Basic 5,000円／月　＋　Premium 10,000円／月",
             size=24, color=WHITE, bold=True)
    add_text(s, Inches(0.6), Inches(6.15), Inches(12.1), Inches(0.4),
             "2026年8月 提供開始予定", size=16, color=WHITE)


# ─────────────────────────────────────────────
# Slide 2: なぜサブスクなのか
# ─────────────────────────────────────────────
def slide_why(prs):
    s = slide_blank(prs)
    add_header_bar(s, "なぜ「サブスク」を始めるのか",
                   "卒業後も、AIの進化に置いていかれないために")

    items = [
        ("🚀", "AI は3ヶ月で景色が変わる",
         "ツールも価格も、毎月どこかが更新されます。\n最新情報を“勝手に”受け取れる場所が必要です。"),
        ("🤝", "一人だと続かない、を解決",
         "Discord と LINE で仲間と運営が伴走。\n質問できる場所があるから、止まらない。"),
        ("💸", "継続するほど収益が積み上がる",
         "動画・楽曲・Kindleは「デジタル資産」。\n運用を止めない人だけが、雪だるま式に伸びます。"),
    ]
    top = Inches(2.2)
    card_w = Inches(4.0)
    card_h = Inches(4.2)
    gap = Inches(0.27)
    left0 = Inches(0.6)
    for i, (icon, title, body) in enumerate(items):
        left = left0 + (card_w + gap) * i
        add_round_rect(s, left, top, card_w, card_h, WHITE, line=BLUE)
        add_rect(s, left, top, card_w, Inches(0.18), BLUE)
        add_text(s, left, top + Inches(0.5), card_w, Inches(0.9), icon,
                 size=44, align=PP_ALIGN.CENTER)
        add_text(s, left + Inches(0.2), top + Inches(1.6),
                 card_w - Inches(0.4), Inches(0.9),
                 title, size=18, color=BLUE_DARK, bold=True, align=PP_ALIGN.CENTER)
        add_text(s, left + Inches(0.3), top + Inches(2.5),
                 card_w - Inches(0.6), Inches(1.6),
                 body, size=14, color=DARK, align=PP_ALIGN.CENTER)


# ─────────────────────────────────────────────
# Slide 3: 2プラン比較（メインスライド）
# ─────────────────────────────────────────────
def slide_compare(prs):
    s = slide_blank(prs)
    add_header_bar(s, "2つのプランから、あなたに合うほうを",
                   "Basic は「最新情報を取り続ける」／ Premium は「直接サポートを受ける」")

    # Basic card
    bx = Inches(0.6)
    by = Inches(2.2)
    bw = Inches(6.0)
    bh = Inches(5.0)
    add_round_rect(s, bx, by, bw, bh, WHITE, line=GRAY)
    add_rect(s, bx, by, bw, Inches(0.7), GRAY)
    add_text(s, bx, by + Inches(0.05), bw, Inches(0.6),
             "Basic ベーシック", size=22, color=WHITE, bold=True,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, bx, by + Inches(0.9), bw, Inches(0.5),
             "月額 5,000円（税込）", size=26, color=DARK, bold=True,
             align=PP_ALIGN.CENTER)
    add_text(s, bx, by + Inches(1.4), bw, Inches(0.35),
             "1日あたり 約167円", size=12, color=GRAY, align=PP_ALIGN.CENTER)

    basic_items = [
        "✅ 最新講義アーカイブ 見放題",
        "✅ Discord コミュニティ参加（無期限）",
        "✅ 仲間とのピア質問・成果報告",
        "✅ 月次の最新AIツール速報",
        "✅ 配布テンプレート 200種＋ 利用可",
    ]
    for i, t in enumerate(basic_items):
        add_text(s, bx + Inches(0.5), by + Inches(1.95) + Inches(0.5) * i,
                 bw - Inches(1.0), Inches(0.45),
                 t, size=15, color=DARK)
    add_text(s, bx + Inches(0.4), by + bh - Inches(0.6), bw - Inches(0.8), Inches(0.5),
             "こんな方に： まずは情報だけ追いつきたい／自走できる方",
             size=12, color=GRAY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    # Premium card（強調）
    px = Inches(6.8)
    py = Inches(2.05)
    pw = Inches(6.0)
    ph = Inches(5.15)
    add_round_rect(s, px, py, pw, ph, WHITE, line=ORANGE)
    add_rect(s, px, py, pw, Inches(0.7), ORANGE)
    add_text(s, px, py + Inches(0.05), pw, Inches(0.6),
             "Premium プレミアム  ★ おすすめ",
             size=22, color=WHITE, bold=True,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, px, py + Inches(0.9), pw, Inches(0.5),
             "月額 10,000円（税込）", size=26, color=ORANGE, bold=True,
             align=PP_ALIGN.CENTER)
    add_text(s, px, py + Inches(1.4), pw, Inches(0.35),
             "1日あたり 約333円（コーヒー1杯分）",
             size=12, color=GRAY, align=PP_ALIGN.CENTER)

    premium_items = [
        "✅ Basic の特典をすべて含む",
        "🌟 月1回 グループコンサル（Zoom 60〜90分）",
        "🌟 Discord 個別質問対応（運営が直接回答）",
        "🌟 LINE ビジネス壁打ちグループへの招待",
        "🌟 メルマガ／引き寄せサミットラインで作品紹介",
    ]
    for i, t in enumerate(premium_items):
        add_text(s, px + Inches(0.5), py + Inches(1.95) + Inches(0.5) * i,
                 pw - Inches(1.0), Inches(0.45),
                 t, size=15, color=DARK, bold=(i > 0))
    add_text(s, px + Inches(0.4), py + ph - Inches(0.6), pw - Inches(0.8), Inches(0.5),
             "こんな方に： 直接フィードバックが欲しい／伴走サポートが欲しい方",
             size=12, color=ORANGE, bold=True,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


# ─────────────────────────────────────────────
# Slide 4: 詳細特典マトリクス
# ─────────────────────────────────────────────
def slide_matrix(prs):
    s = slide_blank(prs)
    add_header_bar(s, "特典 早見表",
                   "Premium だけの特典は ★。それ以外はどちらのプランでも受けられます。")

    rows = [
        ("項目", "Basic", "Premium"),
        ("最新講義アーカイブ", "○", "○"),
        ("Discord コミュニティ参加", "○", "○"),
        ("配布テンプレート 200種＋", "○", "○"),
        ("最新AIツール速報", "○", "○"),
        ("Discord 個別質問への運営回答", "△ ピア中心", "★ 運営が直接対応"),
        ("月1回 グループコンサル（Zoom）", "—", "★ 60〜90分／録画あり"),
        ("LINE ビジネス壁打ちグループ", "—", "★ 招待付き"),
        ("メルマガでの作品告知・紹介", "—", "★ ookuma 担当"),
    ]
    top = Inches(2.15)
    left = Inches(0.6)
    table_w = Inches(12.1)
    col_w = [Inches(5.7), Inches(3.1), Inches(3.3)]
    row_h = Inches(0.55)
    for r, row in enumerate(rows):
        x = left
        is_header = (r == 0)
        bg = BLUE if is_header else (LIGHT_BG if r % 2 == 0 else WHITE)
        for c, cell in enumerate(row):
            add_rect(s, x, top + row_h * r, col_w[c], row_h, bg, line=GRAY)
            color = WHITE if is_header else DARK
            bold = is_header or c == 0
            align = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER
            tb = s.shapes.add_textbox(x + Inches(0.15), top + row_h * r,
                                      col_w[c] - Inches(0.3), row_h)
            tf = tb.text_frame
            tf.word_wrap = True
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf.margin_left = tf.margin_right = Emu(0)
            tf.margin_top = tf.margin_bottom = Emu(0)
            p = tf.paragraphs[0]
            p.alignment = align
            run = p.add_run()
            size = 14 if is_header else 13
            if not is_header and c > 0 and cell.startswith("★"):
                set_run(run, cell, size=size, color=ORANGE, bold=True)
            elif not is_header and c > 0 and cell == "○":
                set_run(run, cell, size=size, color=BLUE_DARK, bold=True)
            else:
                set_run(run, cell, size=size, color=color, bold=bold)
            x += col_w[c]


# ─────────────────────────────────────────────
# Slide 5: Premium 特典の詳細（深掘り）
# ─────────────────────────────────────────────
def slide_premium_deep(prs):
    s = slide_blank(prs)
    add_header_bar(s, "Premium だけの 4 つの特典",
                   "“ひとりで止まる”を、構造的になくします")

    items = [
        ("🎙️", "月1回 グループコンサル",
         "Zoom 60〜90分／月1固定枠。\n事前質問フォーム → 個別回答 → フリー Q&A。\n参加できなくても録画でキャッチアップ。"),
        ("💬", "Discord 個別質問対応",
         "ビジネス系：HERO（桜木先生）が回答。\n占い系：ookuma さんが回答。\n専用チャンネル `#💎premium-質問箱` でやり取り。"),
        ("📱", "LINE ビジネス壁打ちグループ",
         "ビジネス文脈の即時相談チャネル。\n平日24時間以内の応答を目安に運用。\n占い・スピリチュアルとは別運用で混ざりません。"),
        ("📣", "メルマガ・サミットラインで作品紹介",
         "あなたの AI 作品を、ookuma さんが\nメルマガ／引き寄せサミットラインで\n紹介。新しいファンと出会えます。"),
    ]
    top = Inches(2.15)
    card_w = Inches(6.0)
    card_h = Inches(2.45)
    gap_x = Inches(0.1)
    gap_y = Inches(0.2)
    left0 = Inches(0.6)
    for i, (icon, title, body) in enumerate(items):
        col = i % 2
        row = i // 2
        x = left0 + (card_w + gap_x) * col
        y = top + (card_h + gap_y) * row
        add_round_rect(s, x, y, card_w, card_h, WHITE, line=ORANGE)
        add_rect(s, x, y, Inches(0.18), card_h, ORANGE)
        add_text(s, x + Inches(0.3), y + Inches(0.15), Inches(0.9), Inches(0.9),
                 icon, size=36)
        add_text(s, x + Inches(1.2), y + Inches(0.2), card_w - Inches(1.4), Inches(0.5),
                 title, size=18, color=ORANGE, bold=True)
        add_text(s, x + Inches(1.2), y + Inches(0.85), card_w - Inches(1.4),
                 card_h - Inches(1.0),
                 body, size=13, color=DARK)


# ─────────────────────────────────────────────
# Slide 6: 質問対応の分担
# ─────────────────────────────────────────────
def slide_qa_split(prs):
    s = slide_blank(prs)
    add_header_bar(s, "質問の答えは、得意な人から返ります",
                   "ジャンルで担当を分け、迷わずに最短で解決へ")

    cards = [
        ("💼", "ビジネス・マネタイズ\nツール操作・環境構築",
         "HERO（桜木先生）",
         "Discord `#❓質問箱` ／ `#💎premium-質問箱`\nLINE ビジネス壁打ちグループ",
         BLUE_DARK),
        ("🔮", "占い・スピリチュアル\n作品の世界観",
         "ookuma さん",
         "Discord 占い系タグ ／\nメルマガ・引き寄せサミットライン",
         ORANGE),
    ]
    top = Inches(2.2)
    card_w = Inches(6.0)
    card_h = Inches(4.8)
    gap = Inches(0.1)
    left0 = Inches(0.6)
    for i, (icon, theme, who, where, color) in enumerate(cards):
        x = left0 + (card_w + gap) * i
        add_round_rect(s, x, top, card_w, card_h, WHITE, line=color)
        add_rect(s, x, top, card_w, Inches(0.7), color)
        add_text(s, x, top + Inches(0.1), card_w, Inches(0.5),
                 "得意ジャンルで対応します",
                 size=14, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
        add_text(s, x, top + Inches(0.9), card_w, Inches(0.9),
                 icon, size=48, align=PP_ALIGN.CENTER)
        add_text(s, x + Inches(0.4), top + Inches(1.95), card_w - Inches(0.8), Inches(1.0),
                 theme, size=18, color=DARK, bold=True, align=PP_ALIGN.CENTER)
        add_text(s, x, top + Inches(3.0), card_w, Inches(0.5),
                 "→ " + who, size=22, color=color, bold=True, align=PP_ALIGN.CENTER)
        add_text(s, x + Inches(0.4), top + Inches(3.7), card_w - Inches(0.8), Inches(1.0),
                 where, size=12, color=GRAY, align=PP_ALIGN.CENTER)


# ─────────────────────────────────────────────
# Slide 7: 月1コンサル運用
# ─────────────────────────────────────────────
def slide_consult(prs):
    s = slide_blank(prs)
    add_header_bar(s, "Premium 月1グループコンサルの流れ",
                   "事前質問フォーム → 当日回答 → 録画でいつでも復習")

    steps = [
        ("1", "事前", "Discord で質問フォームを公開\n気になることを書くだけ"),
        ("2", "当日 冒頭", "今月の AI トレンド共有\n（10分・HERO）"),
        ("3", "当日 メイン", "個別質問への回答\n（30〜50分）"),
        ("4", "当日 後半", "フリー Q&A／壁打ち\n（残り時間）"),
        ("5", "事後", "録画を `#💎premium-lounge` に投稿\n欠席しても OK"),
    ]
    top = Inches(2.4)
    item_w = Inches(2.3)
    item_h = Inches(3.6)
    gap = Inches(0.15)
    total_w = item_w * 5 + gap * 4
    left0 = (prs.slide_width - total_w) // 2
    for i, (num, label, body) in enumerate(steps):
        x = left0 + (item_w + gap) * i
        add_round_rect(s, x, top, item_w, item_h, WHITE, line=BLUE)
        circ = s.shapes.add_shape(MSO_SHAPE.OVAL, x + item_w / 2 - Inches(0.45),
                                   top - Inches(0.45), Inches(0.9), Inches(0.9))
        circ.fill.solid()
        circ.fill.fore_color.rgb = BLUE
        circ.line.fill.background()
        circ.shadow.inherit = False
        tf = circ.text_frame
        tf.margin_left = tf.margin_right = Emu(0)
        tf.margin_top = tf.margin_bottom = Emu(0)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        set_run(p.add_run(), num, size=24, color=WHITE, bold=True)
        add_text(s, x, top + Inches(0.7), item_w, Inches(0.5),
                 label, size=16, color=BLUE_DARK, bold=True, align=PP_ALIGN.CENTER)
        add_text(s, x + Inches(0.2), top + Inches(1.3), item_w - Inches(0.4),
                 item_h - Inches(1.4),
                 body, size=12, color=DARK, align=PP_ALIGN.CENTER)

    add_rect(s, Inches(0.6), Inches(6.4), Inches(12.1), Inches(0.7), LIGHT_BG)
    add_text(s, Inches(0.8), Inches(6.4), Inches(11.7), Inches(0.7),
             "📅 開催枠（予定）：毎月 第3木曜 20:00〜　／　Zoom リンクは `#💎月1コンサル予告` に掲載",
             size=14, color=DARK, bold=True, anchor=MSO_ANCHOR.MIDDLE)


# ─────────────────────────────────────────────
# Slide 8: 申込・変更・解約
# ─────────────────────────────────────────────
def slide_ops(prs):
    s = slide_blank(prs)
    add_header_bar(s, "申込み・プラン変更・解約について",
                   "わかりやすく、いつでも変更できます")

    blocks = [
        ("📝 お申込み",
         ["決済ページから Basic / Premium を選択",
          "決済完了後、Discord ロールが自動付与",
          "Premium は LINE グループ招待 URL もメール送付"]),
        ("🔄 プラン変更",
         ["Basic ⇄ Premium はマイページからいつでも変更可",
          "アップグレードは即日反映",
          "ダウングレードは次回更新日から反映"]),
        ("🛑 解約",
         ["マイページの解約ボタンから手続き",
          "次回課金日の前日までに手続きで翌月分は不要",
          "解約後も既存卒業生 Discord には残れます"]),
    ]
    top = Inches(2.2)
    card_w = Inches(4.0)
    card_h = Inches(4.4)
    gap = Inches(0.27)
    left0 = Inches(0.6)
    for i, (title, lines) in enumerate(blocks):
        x = left0 + (card_w + gap) * i
        add_round_rect(s, x, top, card_w, card_h, WHITE, line=BLUE)
        add_rect(s, x, top, card_w, Inches(0.7), BLUE)
        add_text(s, x, top + Inches(0.1), card_w, Inches(0.5),
                 title, size=18, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
        for j, line in enumerate(lines):
            add_text(s, x + Inches(0.35), top + Inches(1.0) + Inches(0.85) * j,
                     card_w - Inches(0.7), Inches(0.85),
                     "・" + line, size=14, color=DARK)


# ─────────────────────────────────────────────
# Slide 9: 1日あたりコスト訴求
# ─────────────────────────────────────────────
def slide_cost(prs):
    s = slide_blank(prs)
    add_header_bar(s, "毎月の負担、こう考えると小さく見えてきます",
                   "学びを止めないコストは、収益の積み上げに直結します")

    cards = [
        ("Basic", "5,000円 / 月", "1日 約 167円",
         "缶コーヒー1本分。\n最新情報に毎月触れ続けられる。", BLUE_DARK),
        ("Premium", "10,000円 / 月", "1日 約 333円",
         "コンビニコーヒー1杯分。\nさらに直接サポートとコンサルが付いてくる。", ORANGE),
    ]
    top = Inches(2.2)
    card_w = Inches(5.8)
    card_h = Inches(4.7)
    gap = Inches(0.4)
    left0 = (prs.slide_width - (card_w * 2 + gap)) // 2
    for i, (name, price, daily, body, color) in enumerate(cards):
        x = left0 + (card_w + gap) * i
        add_round_rect(s, x, top, card_w, card_h, WHITE, line=color)
        add_text(s, x, top + Inches(0.4), card_w, Inches(0.5),
                 name, size=22, color=color, bold=True, align=PP_ALIGN.CENTER)
        add_text(s, x, top + Inches(1.0), card_w, Inches(0.7),
                 price, size=32, color=DARK, bold=True, align=PP_ALIGN.CENTER)
        add_rect(s, x + Inches(1.0), top + Inches(2.0), card_w - Inches(2.0),
                 Emu(15000), color)
        add_text(s, x, top + Inches(2.2), card_w, Inches(0.7),
                 daily, size=28, color=color, bold=True, align=PP_ALIGN.CENTER)
        add_text(s, x + Inches(0.5), top + Inches(3.2), card_w - Inches(1.0),
                 Inches(1.4), body,
                 size=14, color=DARK, align=PP_ALIGN.CENTER)


# ─────────────────────────────────────────────
# Slide 10: クロージング
# ─────────────────────────────────────────────
def slide_closing(prs):
    s = slide_blank(prs)
    add_rect(s, 0, 0, prs.slide_width, prs.slide_height, BLUE_DARK)
    add_text(s, Inches(0.6), Inches(1.0), Inches(12.1), Inches(0.6),
             "迷ったら、Premium で 1 ヶ月だけ試す。",
             size=36, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(s, Inches(0.6), Inches(1.9), Inches(12.1), Inches(0.5),
             "合わなければ翌月から Basic にダウングレードできます。",
             size=18, color=WHITE, align=PP_ALIGN.CENTER)

    add_round_rect(s, Inches(2.6), Inches(3.2), Inches(8.1), Inches(2.4),
                   WHITE, line=ORANGE)
    add_text(s, Inches(2.6), Inches(3.4), Inches(8.1), Inches(0.5),
             "✅ 続ける仕組みは、月額にしてしまうのが一番ラク",
             size=18, color=DARK, bold=True, align=PP_ALIGN.CENTER)
    add_text(s, Inches(2.6), Inches(4.0), Inches(8.1), Inches(0.5),
             "✅ AIの最新情報が“勝手に”入ってくる環境を持つ",
             size=18, color=DARK, bold=True, align=PP_ALIGN.CENTER)
    add_text(s, Inches(2.6), Inches(4.6), Inches(8.1), Inches(0.5),
             "✅ 直接質問できる相手がいれば、止まらない",
             size=18, color=DARK, bold=True, align=PP_ALIGN.CENTER)
    add_text(s, Inches(2.6), Inches(5.2), Inches(8.1), Inches(0.5),
             "今日の Discord お知らせから、お申し込みください。",
             size=14, color=ORANGE, bold=True, align=PP_ALIGN.CENTER)

    add_text(s, Inches(0.6), Inches(6.5), Inches(12.1), Inches(0.5),
             "らくらくAI副業キャンパス　|　桜木先生（HERO）／ ookuma",
             size=14, color=WHITE, align=PP_ALIGN.CENTER)


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide_title(prs)
    slide_why(prs)
    slide_compare(prs)
    slide_matrix(prs)
    slide_premium_deep(prs)
    slide_qa_split(prs)
    slide_consult(prs)
    slide_ops(prs)
    slide_cost(prs)
    slide_closing(prs)

    out = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                       "サブスク案内_Basic_Premium_v1.pptx")
    prs.save(out)
    print(f"✅ Saved: {out}  ({len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
