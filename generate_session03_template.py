# -*- coding: utf-8 -*-
"""
第3回 プレゼント用テンプレート集をWordで出力
「AI音声生成プロンプト」「YouTube公開設定プロンプト」「量産バッチ制作テンプレート」3点セット
"""
import sys, os

try:
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
except ImportError:
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx"])
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUT_PATH = os.path.join(SCRIPT_DIR, "第3回_テンプレート3点セット.docx")

# ── ヘルパー ──────────────────────────────────────
def set_font(run, size_pt, bold=False, color=None, name="メイリオ"):
    run.font.name = name
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)

def heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    colors = {1: (0x05, 0x96, 0x69), 2: (0x05, 0x96, 0x69), 3: (0x03, 0x7A, 0x52)}
    sizes  = {1: 20, 2: 15, 3: 13}
    set_font(run, sizes[level], bold=True, color=colors[level])
    return p

def body(doc, text, indent=0, size=11, color=None, bold=False):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(indent)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text)
    set_font(run, size, bold=bold, color=color)
    return p

def box(doc, lines, bg_rgb=(0xF0, 0xFD, 0xF4), border_rgb=(0x05, 0x96, 0x69)):
    for i, line in enumerate(lines):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent  = Cm(0.8)
        p.paragraph_format.right_indent = Cm(0.8)
        p.paragraph_format.space_after  = Pt(1)
        if i == 0:
            p.paragraph_format.space_before = Pt(6)
        if i == len(lines) - 1:
            p.paragraph_format.space_after = Pt(6)
        run = p.add_run(line)
        set_font(run, 11, color=(0x02, 0x4B, 0x2E) if line.startswith("【") else None)
        pPr = p._p.get_or_add_pPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        r, g, b = bg_rgb
        shd.set(qn('w:fill'), f'{r:02X}{g:02X}{b:02X}')
        pPr.append(shd)

def example_box(doc, label, lines):
    p0 = doc.add_paragraph()
    p0.paragraph_format.space_before = Pt(4)
    p0.paragraph_format.space_after  = Pt(2)
    r0 = p0.add_run(f"✏️ 入力例：{label}")
    set_font(r0, 11, bold=True, color=(0x59, 0x39, 0xAB))
    for i, line in enumerate(lines):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent  = Cm(0.8)
        p.paragraph_format.right_indent = Cm(0.8)
        p.paragraph_format.space_after  = Pt(1)
        if i == len(lines) - 1:
            p.paragraph_format.space_after = Pt(8)
        run = p.add_run(line)
        set_font(run, 10.5)
        pPr = p._p.get_or_add_pPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), 'F5F3FF')
        pPr.append(shd)

def tip(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run(f"📌 {text}")
    set_font(run, 10.5, bold=True, color=(0xD9, 0x77, 0x06))

def divider(doc):
    doc.add_paragraph("─" * 55)

def step_box(doc, steps):
    """番号付きステップボックス（薄い青緑背景）"""
    for i, step in enumerate(steps):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent  = Cm(0.8)
        p.paragraph_format.right_indent = Cm(0.8)
        p.paragraph_format.space_after  = Pt(2)
        if i == 0:
            p.paragraph_format.space_before = Pt(6)
        if i == len(steps) - 1:
            p.paragraph_format.space_after = Pt(6)
        run = p.add_run(step)
        bold = step.startswith("STEP") or step.startswith("▶")
        set_font(run, 11, bold=bold,
                 color=(0x02, 0x4B, 0x2E) if bold else None)
        pPr = p._p.get_or_add_pPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), 'E0F2FE')
        pPr.append(shd)

# ── ドキュメント初期化 ────────────────────────────
doc = Document()
sec = doc.sections[0]
sec.left_margin   = Cm(2.0)
sec.right_margin  = Cm(2.0)
sec.top_margin    = Cm(2.0)
sec.bottom_margin = Cm(2.0)

style = doc.styles['Normal']
style.font.name = 'メイリオ'
style.font.size = Pt(11)
style.paragraph_format.line_spacing = 1.4

# ══ 表紙 ══════════════════════════════════════════
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(40)
r = p.add_run("らくらくAI副業キャンパス 第3回\n配布テンプレート")
set_font(r, 18, bold=True, color=(0x05, 0x96, 0x69))

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("プロンプト 3点セット")
set_font(r, 26, bold=True, color=(0x06, 0x4E, 0x3B))

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("① AI音声生成プロンプト　② YouTube公開設定プロンプト　③ 量産バッチ制作テンプレート")
set_font(r, 13, color=(0x37, 0x47, 0x51))

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("\n今日から世界に届ける動画を、AIと一緒に量産しましょう！")
set_font(r, 12, bold=True, color=(0xD9, 0x77, 0x06))

doc.add_page_break()

# ══ 使い方ガイド ══════════════════════════════════
heading(doc, "📖 このテンプレートの使い方", level=1)
body(doc, "1. 使いたいプロンプトをそのままコピーする")
body(doc, "2.「〇〇」の部分を自分のジャンル・テーマ・動画内容に書き換える")
body(doc, "3. Google AI Studio（音声）または Gemini（テキスト）に貼り付けてEnterを押すだけ！")
doc.add_paragraph()
body(doc, "▶ 今日完成した動画のURLを、必ずDiscordに投稿しましょう！仲間みんなで見に行きます。", bold=True)
body(doc, "▶ 量産テンプレートは「週末に7本仕込む」ルーティンに活用してください。")

divider(doc)
doc.add_paragraph()

# ══════════════════════════════════════════════════
# ① AI音声生成プロンプト
# ══════════════════════════════════════════════════
heading(doc, "① AI音声生成プロンプト", level=1)
body(doc, "Google AI Studio（aistudio.google.com）で使います　※ Googleアカウントで無料利用OK",
     color=(0x55, 0x65, 0x74))

heading(doc, "▼ Google AI Studio 音声生成の手順", level=2)
step_box(doc, [
    "STEP 1：aistudio.google.com にアクセス → Googleアカウントでログイン",
    "STEP 2：左メニューから「Speech」を選択",
    "STEP 3：先週Geminiで作った台本テキストをコピペ",
    "STEP 4：声の種類・速さを選んで生成ボタンを押す → mp3でダウンロード",
])

heading(doc, "▼ 声のクオリティを上げる3つのコツ", level=2)
body(doc, "コツ① 速さは 0.85〜0.9倍速 に設定する（デフォルトは少し速すぎる）")
body(doc, "コツ② 台本の句読点「、。」を増やすと、プロのナレーターのような自然な間が生まれる")
body(doc, "コツ③ キャラクターの雰囲気に合った声を選ぶ（神秘的→落ち着いた声、かわいい→明るい声）")
doc.add_paragraph()

heading(doc, "▼ 台本テキストを音声向けに調整するプロンプト（Gemini用）", level=2)
body(doc, "台本はできているけど「音声にすると聞きづらい」と感じたときに使ってください。",
     color=(0x55, 0x65, 0x74))

box(doc, [
    "以下の台本テキストを、AI音声読み上げに適した形に調整してください。",
    "",
    "【調整内容】",
    "・句読点「、。」を自然な間（ま）ができるように適切に追加する",
    "・難しい漢字や専門用語には読み仮名を（ふりがな）でつける",
    "・数字は「1」ではなく「いち」のように読み仮名を入れる",
    "・リズムよく聞こえるよう、長い文は短く区切る",
    "・感嘆符「！」は「。」に変換する（音声では不自然になるため）",
    "",
    "【台本テキスト】",
    "（ここに台本テキストを貼り付けてください）",
])

tip(doc, "調整後のテキストをGoogle AI Studioに貼り付けると、格段に聞きやすい音声になります！")

heading(doc, "▼ ジャンル別 おすすめ声の選び方", level=2)

t = doc.add_table(rows=5, cols=2)
t.style = 'Light Grid Accent 6'
headers = ["チャンネルのジャンル", "おすすめの声の雰囲気"]
for i, h in enumerate(headers):
    c = t.rows[0].cells[i]
    c.text = h
    for p in c.paragraphs:
        for r in p.runs:
            r.font.bold = True
            r.font.name = 'メイリオ'
rows_data = [
    ["スピリチュアル・瞑想系", "落ち着いた低め、ゆっくり、女性か中性的な声"],
    ["猫・ペット・癒やし系", "温かみのある声、ほんの少し高め、親しみやすい"],
    ["雑学・教養系", "はっきりした声、メリハリがある、標準的な速さ"],
    ["子ども向け・キッズ系", "明るく元気な声、少し高め、ゆっくりはっきり"],
]
for ri, row_data in enumerate(rows_data, 1):
    for ci, val in enumerate(row_data):
        c = t.rows[ri].cells[ci]
        c.text = val
        for p in c.paragraphs:
            for r in p.runs:
                r.font.name = 'メイリオ'
                r.font.size = Pt(10)

doc.add_paragraph()
divider(doc)
doc.add_paragraph()

# ══════════════════════════════════════════════════
# ② YouTube公開設定プロンプト
# ══════════════════════════════════════════════════
heading(doc, "② YouTube公開設定プロンプト", level=1)
body(doc, "Gemini（gemini.google.com）に貼り付けて使います", color=(0x55, 0x65, 0x74))
body(doc, "タイトル・説明文・タグの3点は、すべてAIに作ってもらいましょう！")
doc.add_paragraph()

heading(doc, "▼ タイトル生成プロンプト（コピペ用）", level=2)
box(doc, [
    "以下の条件でYouTube動画のタイトルを5個提案してください。",
    "",
    "【動画のジャンル】〇〇（例：スピリチュアル / 猫の雑学 / 歴史）",
    "【動画のテーマ】〇〇（例：直感を磨く3つの習慣 / 猫が甘えてくる理由）",
    "【条件】",
    "  ・30文字以内で収める",
    "  ・クリックしたくなるような言葉を使う",
    "  ・以下のパターンを活用して、バリエーションをつける",
    "    ❶「〇〇な人の特徴」「知らないと損」「実は〇〇だった」",
    "    ❷「【衝撃】〇〇」「なぜ〇〇？」「〇〇の真実」",
    "    ❸「〇〇するだけで」「〇〇してみた」「〇〇選」",
])

heading(doc, "▼ タイトル入力例", level=2)

example_box(doc, "スピリチュアル系", [
    "【動画のジャンル】スピリチュアル",
    "【動画のテーマ】直感を磨くための3つの習慣",
    "【条件】30文字以内・クリックしたくなるタイトル・5個提案",
])

example_box(doc, "猫・癒やし系", [
    "【動画のジャンル】猫の雑学",
    "【動画のテーマ】猫が「ふみふみ」をする本当の理由",
    "【条件】30文字以内・クリックしたくなるタイトル・5個提案",
])

tip(doc, "5個提案してもらって、一番気に入ったものを選ぼう。「もっと短くして」「もっとインパクトを強くして」と追加指示もOK！")

heading(doc, "▼ 説明文生成プロンプト（コピペ用）", level=2)
box(doc, [
    "以下の条件でYouTube動画の説明文を書いてください。",
    "",
    "【動画のジャンル】〇〇",
    "【動画のタイトル】〇〇",
    "【動画の内容（台本の要約）】",
    "（ここに台本の主な内容を2〜3行で書いてください）",
    "",
    "【条件】",
    "  ・冒頭の2行で「この動画で何がわかるか」を伝える（検索結果に表示される部分）",
    "  ・全体で200〜300字程度",
    "  ・自然に「〇〇（キーワード）」を2〜3回盛り込む",
    "  ・最後にチャンネル登録を促す一文を入れる",
    "  ・絵文字を適度に使って読みやすくする",
])

heading(doc, "▼ タグ生成プロンプト（コピペ用）", level=2)
box(doc, [
    "以下の動画に合う YouTube タグを15個提案してください。",
    "",
    "【動画のジャンル】〇〇",
    "【動画のテーマ】〇〇",
    "",
    "【条件】",
    "  ・検索されやすい具体的なキーワードで",
    "  ・短いもの（1〜2語）と長いもの（3〜5語）を混ぜる",
    "  ・競合が少なそうなニッチなキーワードも入れる",
    "  ・カンマ区切りで出力する",
])

tip(doc, "タイトル→説明文→タグの順番で作ると、キーワードが統一されてSEO効果がアップします！")

heading(doc, "▼ サムネイル制作メモ", level=2)
body(doc, "Canva（canva.com）で無料作成。この3点だけ意識すればOK！")
body(doc, "  ✅ 大きな文字（タイトルの核心部分）")
body(doc, "  ✅ 鮮やかな色（背景とキャラクターのコントラストをつける）")
body(doc, "  ✅ シンプルなデザイン（ごちゃごちゃしない）")
body(doc, "サムネイルの良し悪しで再生数が3〜10倍変わります。手を抜かないことが最大のコツです。",
     bold=True, color=(0x02, 0x4B, 0x2E))

doc.add_paragraph()
divider(doc)
doc.add_paragraph()

# ══════════════════════════════════════════════════
# ③ 量産バッチ制作テンプレート
# ══════════════════════════════════════════════════
heading(doc, "③ 量産バッチ制作テンプレート", level=1)
body(doc, "週末に7本まとめて仕込む「バッチ制作」の流れです。毎日作業しなくてOK！",
     color=(0x55, 0x65, 0x74))
doc.add_paragraph()

heading(doc, "▼ 週7本量産プロンプト（ネタ出し用）", level=2)
body(doc, "まずGeminiでこの週の「ネタ7本」を一気に出してもらいます。",
     color=(0x55, 0x65, 0x74))

box(doc, [
    "以下の条件で、YouTube動画のネタ（テーマ）を7個提案してください。",
    "",
    "【チャンネルのジャンル】〇〇（例：スピリチュアル / 猫の雑学）",
    "【ターゲット視聴者】〇〇（例：30〜50代の女性、癒やしを求めている人）",
    "【これまでに作った動画テーマ】",
    "（過去に作ったテーマをここに箇条書きで貼り付けてください）",
    "",
    "【条件】",
    "  ・視聴者が「思わずクリックしたくなる」テーマにする",
    "  ・すでに作ったテーマと被らないようにする",
    "  ・リスト系（「〇〇な人の5つの特徴」）と教養系（「なぜ〇〇？」）を混ぜる",
    "  ・テーマ名と「一行説明」をセットで出力する",
])

tip(doc, "ネタ7個が決まったら、台本プロンプトに1本ずつ入れて台本を一気に7本生成！")

heading(doc, "▼ 台本一括生成プロンプト（1本分・第2回のプロンプトと組み合わせて使う）", level=2)
box(doc, [
    "以下の条件でYouTube動画の台本を書いてください。",
    "",
    "【ジャンル】〇〇",
    "【テーマ】〇〇（ネタ一覧から選んだテーマを入れる）",
    "【台本の型】リスト系（「〇〇な人の5つの特徴」のような構成）",
    "【文字数】2,000字程度",
    "【口調】落ち着いた、優しい語り口調。難しい言葉は使わない",
    "【構成】",
    "  ① 冒頭のつかみ（最初の30秒で興味を引く）",
    "  ② 本編（3〜5項目のリスト形式）",
    "  ③ まとめ（温かい言葉で締める）",
    "  ④ 次も見たくなる一言（チャンネル登録を促す）",
    "",
    "※ 台本が完成したら、このテーマのタイトル案も5個提案してください。",
])

heading(doc, "▼ 週間バッチ制作ルーティン表", level=2)
body(doc, "「毎日やらなくていい」のがバッチ制作の最大のメリットです！",
     bold=True, color=(0x02, 0x4B, 0x2E))
doc.add_paragraph()

t2 = doc.add_table(rows=6, cols=3)
t2.style = 'Light Grid Accent 6'
headers2 = ["曜日", "作業内容", "目安時間"]
for i, h in enumerate(headers2):
    c = t2.rows[0].cells[i]
    c.text = h
    for p in c.paragraphs:
        for r in p.runs:
            r.font.bold = True
            r.font.name = 'メイリオ'
rows_data2 = [
    ["土曜（午前）", "ネタ出し7本 → 台本7本を一括生成", "約60分"],
    ["土曜（午後）", "AI音声7本を生成 → ダウンロード", "約60分"],
    ["日曜（午前）", "動画編集7本（KLING AI or CapCut）", "約90分"],
    ["日曜（午後）", "タイトル・説明文・タグ生成 → YouTube予約投稿", "約60分"],
    ["平日（毎日）", "コメント返信・Discordに進捗投稿・仲間の動画を視聴", "約15分"],
]
for ri, row_data in enumerate(rows_data2, 1):
    for ci, val in enumerate(row_data):
        c = t2.rows[ri].cells[ci]
        c.text = val
        for p in c.paragraphs:
            for r in p.runs:
                r.font.name = 'メイリオ'
                r.font.size = Pt(10)

doc.add_paragraph()
tip(doc, "週末4〜5時間で7本分の素材が揃う！あとは「予約投稿」で毎日1本ずつ自動公開されます。")

heading(doc, "▼ 1本あたりの制作チェックリスト", level=2)
body(doc, "動画1本を完成させるたびに、このチェックリストを使いましょう。", color=(0x55, 0x65, 0x74))

box(doc, [
    "【台本】",
    "  □ Geminiで台本2,000字を生成した",
    "  □ 音声向けに句読点・読み仮名を調整した",
    "",
    "【音声】",
    "  □ Google AI Studioで音声を生成した",
    "  □ 速さを0.85〜0.9倍速に調整した",
    "  □ mp3をダウンロードした",
    "",
    "【動画】",
    "  □ 【プレミアム】KLING AIでリップシンク動画を生成した",
    "  □ 【コスト重視】CapCutで背景画像＋音声＋テロップを合成した",
    "  □ BGMを追加した（CapCut内の著作権フリー音楽）",
    "",
    "【YouTube公開】",
    "  □ Geminiでタイトル5案を生成 → 1つ選んだ",
    "  □ Geminiで説明文を生成した",
    "  □ Geminiでタグ15個を生成した",
    "  □ Canvaでサムネイルを作った",
    "  □ YouTubeにアップ → 公開（or予約投稿）した",
    "",
    "【Discordに投稿】",
    "  □ 公開URLをDiscordに投稿した ← これが一番大事！",
], bg_rgb=(0xFE, 0xF9, 0xE7))

doc.add_paragraph()
divider(doc)
doc.add_paragraph()

# ══ まとめ ══
heading(doc, "📋 まとめ｜3点セットの活用フロー", level=1)
body(doc, "STEP 1：③量産バッチ → Geminiでネタ7本を一気に出す", bold=True)
body(doc, "        ネタが決まったら台本7本を一気に生成", indent=1)
doc.add_paragraph()
body(doc, "STEP 2：①AI音声生成 → Google AI Studioで台本を音声に変換", bold=True)
body(doc, "        速さ・声の種類を調整してmp3でダウンロード", indent=1)
doc.add_paragraph()
body(doc, "STEP 3：動画編集 → KLING AI（プレミアム）or CapCut（コスト重視）", bold=True)
body(doc, "        BGM・テロップを追加して動画を完成させる", indent=1)
doc.add_paragraph()
body(doc, "STEP 4：②YouTube公開設定 → タイトル・説明文・タグをGeminiで生成", bold=True)
body(doc, "        サムネイルをCanvaで作って、予約投稿で毎日1本ずつ公開！", indent=1)
doc.add_paragraph()
body(doc, "📌 公開URLを必ずDiscordに投稿！仲間みんなで応援し合いましょう！",
     bold=True, color=(0xD9, 0x77, 0x06))

doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("らくらくAI副業キャンパス | 第3回配布資料")
set_font(r, 10, color=(0x9C, 0xA3, 0xAF))

doc.save(OUT_PATH)
print(f"完了！ → {OUT_PATH}")
