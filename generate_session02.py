# -*- coding: utf-8 -*-
"""第2回「AI動画① チャンネル開設・台本・AIキャラクター誕生」改訂版"""
import sys, os
sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "Google_Slide_generator"))
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from design_v2 import *
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
L_C1, L_C2 = "6366F1", "8B5CF6"

def build(prs):
    # ─── タイトル ───
    s = add_slide(prs)
    make_title_slide(s, "AI動画①\nチャンネル開設・台本・AIキャラクター誕生", "第2回｜2026年4月27日（月）19:00〜21:00", "らくらくAI副業キャンパス")

    # ─── ゴール ───
    s = add_slide(prs)
    make_cards_slide(s, "🎯 本日のゴール", [
        ("📺", "YouTubeチャンネル\n土台作り", "今日中にチャンネルを\n作って公開準備完了\n（ゆっくりOK！）"),
        ("📝", "台本1本を\n完成させる", "Geminiと会話しながら\nあなたらしい台本を生成"),
        ("🎭", "AIキャラクターを\n誕生させる", "あなたの世界観を体現する\n動画の「顔」を\n今日作ります！"),
    ], accent_c1=L_C1, accent_c2=L_C2)

    # ─── 振り返り ───
    s = add_slide(prs)
    make_dark_summary(s, "前回の振り返り", [
        "✅ Geminiの基本操作をマスター",
        "✅ プロンプト黄金の3パターン（指示・質問・条件）",
        "✅ Discordで温かい仲間と繋がった",
    ], sub="今日から「あなたの世界観」を形にするフェーズです", c1=L_C1, c2=L_C2)

    # ─── 課題チェック ───
    s = add_slide(prs)
    make_cards_slide(s, "課題チェック｜Geminiと仲良くなれましたか？", [
        ("💬", "10回以上\n話しかけた？", "面白かった回答を\nピックアップ紹介します！"),
        ("📸", "Discordに\n投稿できた？", "ベスト回答を\n画面で共有します"),
        ("👍", "仲間に\nリアクションした？", "コミュニティが\n温かい雰囲気です！"),
    ], accent_c1=L_C1, accent_c2=L_C2)

    # ══════════════════
    # Section 1：YouTubeの魅力 ＋ WOW体験
    # ══════════════════
    s = add_slide(prs)
    make_section(s, 1, "YouTubeで発信する楽しさ\n〜 AIが「普通では無理」を可能にします 〜", "20:00 - 35:00", c1=L_C1, c2=L_C2)

    s = add_slide(prs)
    make_cards_slide(s, "YouTube発信の3つの魅力", [
        ("🌱", "自動で働く分身", "あなたが温泉旅行中も\n動画が代わりに\nあなたの想いを伝えます"),
        ("🏠", "デジタル資産", "一度アップした動画は\nずっと残る大切な\n「作品」になります"),
        ("🎭", "AIキャラクターが\n活躍する時代", "顔出し不要！\n自分だけのAIアバターが\n世界中に発信します"),
    ], accent_c1=L_C1, accent_c2=L_C2)

    s = add_slide(prs)
    make_emphasis(s, "でも、普通の動画作りって\n「すごく大変」ですよね…？",
        "・高い機材・編集スキルが必要\n・顔出しに抵抗がある\n・台本を考えるだけで何日も悩む\n・動画編集ソフトが難しくて挫折する\n\n→ だから多くの人は「見るだけ」で終わってしまいます", gold=False)

    s = add_slide(prs)
    make_emphasis(s, "そこで「AI」の魔法を使います！\n今日、皆さんは魔法使いになります",
        "たった1行の「思いつき」が、AIの力で数秒にして\n「プロ級の台本」「美しい背景画像」「あなたの分身キャラクター」に変わる。\n\n信じられないかもしれませんが、本当にお見せします！", gold=True)

    s = add_slide(prs)
    make_cards_slide(s, "✨ 魔法のビフォーアフター（台本編）", [
        ("❌", "ただの1行", "「猫と暮らす幸せな\n日常の台本を書いて」"),
        ("🪄", "AIの魔法\n（待ち時間：5秒）", ""),
        ("✅", "涙を誘う名作", "「その小さな温もりが、\n私の孤独を救ってくれた…」\nプロ顔負けの台本完成！"),
    ], accent_c1=L_C1, accent_c2=L_C2)

    s = add_slide(prs)
    make_cards_slide(s, "✨ 魔法のビフォーアフター（キャラクター編）", [
        ("❌", "ただの1行", "「癒やし系の\n女性キャラクターを\n生成して」"),
        ("🪄", "AIの魔法\n（待ち時間：3秒）", ""),
        ("✅", "あなたの分身誕生！", "個性的な表情・衣装の\nオリジナルキャラクターが\n一瞬で完成！"),
    ], accent_c1=L_C1, accent_c2=L_C2)

    s = add_slide(prs)
    make_table_slide(s, "こんな素敵なAIキャラクターチャンネルが！",
        ["チャンネル名", "キャラクター", "ジャンル", "魅力"],
        [
            ["スピリチュアルCh", "神秘的な女性AI", "占い・スピ", "癒やしの言葉でファン多数"],
            ["大人の雑学", "学者風AIアバター", "教育", "趣味の知識が誰かの役に立つ"],
            ["心のお悩み相談", "優しいAI先生", "癒やし", "コメント欄が温かい居場所に"],
            ["ペットの日常", "ゆるキャラAI", "趣味", "楽しく作れてお小遣いにも"],
        ],
        accent_c1=L_C1, accent_c2=L_C2,
        note="📌 すべて顔出しなし・AIキャラクターが主役のチャンネルです")

    s = add_slide(prs)
    make_emphasis(s, "ジャンル選びのコツ\n「無理なく続く × 初心者向き × 自分の好きなこと」",
        "特別なスキルは一切不要です。「なんとなく好き」「ちょっと興味がある」で十分！\nGeminiに「この3条件でYouTubeジャンルを提案して」と聞くだけ！", gold=True)

    # ══════════════════
    # Section 2：チャンネル開設 ＆ 台本生成
    # ══════════════════
    s = add_slide(prs)
    make_section(s, 2, "チャンネル開設 ＆ 台本生成\n〜 Geminiが優秀な秘書になります 〜", "35:00 - 55:00", c1=L_C1, c2=L_C2)

    s = add_slide(prs)
    make_emphasis(s, "PC操作が苦手でも大丈夫！",
        "最初は誰でも時間がかかります。周りと比べず、焦らず自分のペースで。\n\n「できなくてもあなたのせいじゃありません」\n私たちがしっかりサポートします！", gold=True)

    s = add_slide(prs)
    make_step_flow(s, "YouTubeチャンネルの土台作り", [
        ("1", "YouTubeを開く", "youtube.com に\nGoogleアカウントで\nログイン"),
        ("2", "チャンネル作成", "右上アイコン →\n「チャンネルを作成」\nをクリック"),
        ("3", "チャンネル名を\n決める", "Geminiに\n「チャンネル名を\n10個提案して」"),
        ("4", "アイコン設定", "今日作る\nAIキャラクターで\n設定します！"),
    ], accent_c1=L_C1, accent_c2=L_C2)

    s = add_slide(prs)
    make_numbered_slide(s, "台本の型｜3パターン", [
        ("1", "教養系",
         "「知って得する〇〇の真実」「〇〇のランキングTOP10」\n→ 誰でも作りやすく、たくさんの人に見られやすいです"),
        ("2", "物語系",
         "「〇〇した結果、衝撃の結末が…」「〇〇の不思議な話」\n→ 最後まで見てもらいやすい。怖い話・スピ系に最適です"),
        ("3", "リスト系",
         "「〇〇な人の5つの特徴」「やってはいけない〇〇3選」\n→ 構成が簡単で一番作りやすい。初心者に最おすすめ！"),
    ], accent_c1=L_C1, accent_c2=L_C2)

    s = add_slide(prs)
    make_cards_slide(s, "【解説】なぜ「リスト系」が初心者に最適？", [
        ("💡", "人間の心理", "「3つの特徴」と言われると\n人間は全部聞きたくなる\n心理があります"),
        ("🤖", "AIが得意", "AIは箇条書きで\n整理して出力するのが\n非常に得意です"),
        ("✂️", "編集がラク", "「その１」「その２」と\n動画を区切りやすく\n後の作業が超カンタン！"),
    ], accent_c1=L_C1, accent_c2=L_C2)

    s = add_slide(prs)
    bg_grad(s, DARK1, "162544")
    txt(s, Inches(1), Inches(0.5), Inches(11), Inches(0.7),
        "台本生成プロンプト例（Geminiにコピペ）", sz=26, clr=C_GOLD, bold=True)
    card(s, Inches(1.2), Inches(1.4), Inches(10.9), Inches(4.9),
         grad=(DARK2, "2D3748"), radius=0.06, shadow=True, border=RGBColor(0x63, 0x66, 0xF1))
    multi_txt(s, Inches(1.8), Inches(1.7), Inches(10), Inches(4.4), [
        "以下の条件でYouTube動画の台本を書いてください。",
        "",
        "【条件】",
        "・ジャンル：〇〇（例：スピリチュアル・猫・雑学など）",
        "・テーマ：〇〇（例：運気が上がる朝の習慣）",
        "・台本の型：リスト系 / 教養系 / 物語系",
        "・文字数：2,000字程度",
        "・口調：落ち着いた、優しい語り口調",
        "・構成：冒頭のつかみ → 本編 → まとめ → 次も見たくなる一言",
    ], sz=15, clr=C_WHITE, align=PP_ALIGN.LEFT)
    txt(s, Inches(1.2), Inches(6.5), Inches(10.9), Inches(0.5),
        "📌 専門用語不要！コピペして〇〇を変えるだけで素敵な台本が完成します", sz=13, clr=C_GOLD)
    footer(s, L_C1, L_C2)

    s = add_slide(prs)
    make_emphasis(s, "【解説】AIにお願いするときの「最大のコツ」",
        "AIは「とても優秀だけど、ちょっとドジな秘書」だと思ってください。\n\n一回で完璧なものが出てこなくても大丈夫。\n「もう少し短くして！」「もっと優しい言葉遣いにして！」と\n会話しながら直してもらうのがコツです。", gold=True)

    # ══════════════════
    # Section 3：AI画像生成（背景 ＋ キャラクター）
    # ══════════════════
    s = add_slide(prs)
    make_section(s, 3, "AI画像生成\n〜 背景 ＋ あなたのキャラクターを作ろう 〜", "55:00 - 75:00", c1=L_C1, c2=L_C2)

    s = add_slide(prs)
    make_cards_slide(s, "今日作る画像は2種類", [
        ("🖼️", "背景画像", "動画の「舞台」になる画像\n世界観を表す風景や\n幻想的なシーンを生成"),
        ("🎭", "AIキャラクター", "動画の「顔」になる\nあなたの分身キャラ\n来週、ここに命を吹き込みます"),
        ("💡", "使うツール", "Gemini（無料）か\nCanva AI（無料）で\n全部できます！"),
    ], accent_c1=L_C1, accent_c2=L_C2)

    s = add_slide(prs)
    bg_grad(s, DARK1, "162544")
    txt(s, Inches(1), Inches(0.5), Inches(11), Inches(0.7),
        "背景画像 生成プロンプト例", sz=26, clr=C_GOLD, bold=True)
    card(s, Inches(1.2), Inches(1.4), Inches(10.9), Inches(3.0),
         grad=(DARK2, "2D3748"), radius=0.06, shadow=True, border=RGBColor(0x63, 0x66, 0xF1))
    multi_txt(s, Inches(1.8), Inches(1.7), Inches(10), Inches(2.5), [
        "「〇〇のテーマで、YouTube動画の背景画像を作って。",
        "　文字は入れないで。横長サイズで。",
        "　スタイル：水彩画風 / イラスト風 / 幻想的",
        "　雰囲気：神秘的 / 温かい / 落ち着いた」",
    ], sz=15, clr=C_WHITE, align=PP_ALIGN.LEFT)

    txt(s, Inches(1), Inches(4.6), Inches(11), Inches(0.5),
        "AIキャラクター 生成プロンプト例", sz=22, clr=C_GOLD, bold=True)
    card(s, Inches(1.2), Inches(5.2), Inches(10.9), Inches(2.1),
         grad=(DARK2, "2D3748"), radius=0.06, shadow=True, border=RGBColor(0x63, 0x66, 0xF1))
    multi_txt(s, Inches(1.8), Inches(5.4), Inches(10), Inches(1.7), [
        "「〇〇な雰囲気の女性キャラクターを作って。背景は透過（白背景）で。",
        "　スタイル：アニメ風 / リアル系 / ゆるいイラスト風",
        "　特徴：優しい表情・〇〇な衣装・〇〇な髪色」",
    ], sz=15, clr=C_WHITE, align=PP_ALIGN.LEFT)
    txt(s, Inches(1.2), Inches(7.2), Inches(10.9), Inches(0.4),
        "📌 背景とキャラは「雰囲気を揃えた言葉」で統一感を出すのがコツです", sz=12, clr=C_GOLD)
    footer(s, L_C1, L_C2)

    s = add_slide(prs)
    make_emphasis(s, "【解説】画像生成AIは「素直すぎる画伯」",
        "「女性キャラを作って」だけでは、どんな女性か迷ってしまいます。\n\n「神秘的な雰囲気の、長い黒髪の女性。アニメ風。優しい微笑み。白背景。」\nこのように情景を具体的に伝えると、劇的にクオリティが上がります！\n\n📌 背景とキャラは同じ「世界観の言葉」を使うと統一感が出ます", gold=False)

    s = add_slide(prs)
    make_cards_slide(s, "【解説】キャラクター画像のポイント3つ", [
        ("⬜", "背景は白で\n生成する", "「背景透過」か\n「白背景」で生成すると\n後で使いやすい"),
        ("😊", "表情は\nシンプルに", "最初は「優しい微笑み」が\n万能。複雑な表情は\n慣れてから挑戦"),
        ("🎨", "台本の世界観と\n合わせる", "スピ系なら神秘的に\n雑学系ならカジュアルに\n世界観を統一しよう"),
    ], accent_c1=L_C1, accent_c2=L_C2)

    s = add_slide(prs)
    make_emphasis(s, "📌 量産テンプレート配布",
        "台本生成プロンプト ＋ 背景画像プロンプト ＋ キャラクタープロンプト\nの3点セットをテンプレートとしてお渡しします\n「〇〇の部分を変えるだけで、何本でも作れます」", gold=True)

    # ══════════════════
    # Section 4：ハンズオン
    # ══════════════════
    s = add_slide(prs)
    make_section(s, 4, "ハンズオン実践\n〜 一緒に魔法を使ってみよう 〜", "75:00 - 105:00", c1=L_C1, c2=L_C2)

    s = add_slide(prs)
    make_numbered_slide(s, "✋ ハンズオン — 今日のワーク", [
        ("1", "YouTubeチャンネルの土台を作る",
         "画面を見ながらゆっくり進めます。つまずいたらすぐ教えてください"),
        ("2", "ジャンルを確定してGeminiで台本1本を生成する",
         "テンプレートを使って、ワクワクする台本を完成させましょう"),
        ("3", "背景画像を1枚生成する",
         "あなたの世界観を表す「舞台」を作りましょう"),
        ("4", "AIキャラクターを1体生成する",
         "あなたの動画の「顔」となる分身キャラを誕生させましょう！"),
    ], accent_c1=L_C1, accent_c2=L_C2)

    s = add_slide(prs)
    make_emphasis(s, "失敗しても大丈夫。AIは怒りません",
        "何度やり直しても、AIは疲れませんし、怒りません。\n「なんか違うな」と思ったら「もう一回書いて！」と気軽にやり直してください。\n\nキャラクターが思ったのと違っても全然OK。\n「もう少し優しい表情に」「髪を長くして」と会話するのが正解です", gold=True)

    # ─── Closing ───
    s = add_slide(prs)
    make_numbered_slide(s, "📝 今週の課題", [
        ("1", "台本と背景画像・キャラクターを仕上げる",
         "ハンズオンで作ったものを納得いくまで仕上げましょう。完璧じゃなくていいです"),
        ("2", "動画のタイトルをDiscordに投稿",
         "「こんな動画を作ります！」と宣言しましょう。仲間が応援してくれます"),
        ("3", "仲間の投稿に「見たい！」とリアクション",
         "お互いに応援し合うのがキャンパスのルールです"),
    ], accent_c1=L_C1, accent_c2=L_C2)

    s = add_slide(prs)
    make_table_slide(s, "📊 収益源ダッシュボード — 第2回時点",
        ["収益源", "ステータス", "メモ"],
        [
            ["① YouTube動画", "🟡 準備中", "チャンネル開設・台本・キャラ完成"],
            ["② AI音楽", "⬜ まだ", "第5回で開始"],
            ["③ Kindle出版", "⬜ まだ", "第6回で開始"],
            ["④ AIブログ", "⬜ まだ", "第7回で開始"],
            ["⑤ 画像販売", "⬜ まだ", "第8回で開始"],
            ["⑥ ショート動画", "⬜ まだ", "第9回で開始"],
            ["⑦ スキル販売", "⬜ まだ", "第10回で開始"],
        ],
        accent_c1=L_C1, accent_c2=L_C2, highlight_last=False,
        note="📌 来週、あなたのキャラクターが喋る動画を世界に届けます！")

    s = add_slide(prs)
    make_emphasis(s, "🎬 来週の予告\n「AI動画② キャラクターが喋る！世界に届ける日」",
        "来週は今日作ったキャラクターに「声」をつけます。\nGoogle AI Studioで音声 → KLING AIでリップシンク →YouTubeに公開！\n📌 今週の課題：台本・背景画像・キャラクターを仕上げてきてください", gold=True)

    s = add_slide(prs)
    make_emphasis(s, "お疲れさまでした！\nあなたのAIキャラクターが生まれた日です 🎭",
        "📌 課題：台本・背景画像・キャラクター画像の完成\n📌 タイトルをDiscordに投稿して仲間に宣言\n📌 来週月曜19:00にお会いしましょう！")

if __name__ == "__main__":
    prs = create_prs()
    print("第2回を生成中...")
    build(prs)
    out = os.path.join(SCRIPT_DIR, "講義02_AI動画①_v3.pptx")
    prs.save(out)
    print(f"完了！ スライド{len(prs.slides)}枚 → {out}")
